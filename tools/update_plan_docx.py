from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
source = root / "GYM_Management_Master_Plan_V1.3.docx"
target = root / "GYM_Management_Master_Plan_V1.4.docx"
doc = Document(source)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.startswith("V1.3 – Đã cập nhật"):
                cell.text = "V1.4 – Bổ sung tài khoản Owner mặc định và quyết định triển khai MVP"

inserted = False
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "03. DANH MỤC PHÂN HỆ VÀ ƯU TIÊN TRIỂN KHAI":
        added = paragraph.insert_paragraph_before(
            "Tài khoản Owner (chủ phòng) mặc định của dự án: vhoanglong54@gmail.com. "
            "Mật khẩu khởi tạo lấy từ biến môi trường SEED_PASSWORD; không ghi mật khẩu thật trong mã nguồn hoặc tài liệu."
        )
        added.style = doc.styles["Normal"]
        inserted = True
        break

if not inserted:
    raise RuntimeError("Không tìm thấy vị trí Mục 03 để chèn tài khoản Owner.")

doc.core_properties.title = "GYM Management System – Master Plan V1.4"
doc.core_properties.subject = "Bổ sung tài khoản Owner và quyết định triển khai MVP"
doc.save(target)
print(target)
